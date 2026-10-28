from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import usg_parser

_HEADER_FILL = PatternFill(start_color="2F5BEA", end_color="2F5BEA", fill_type="solid")
_DENY_FILL = PatternFill(start_color="FDECEC", end_color="FDECEC", fill_type="solid")
_NAT_FILL = PatternFill(start_color="EEF3FF", end_color="EEF3FF", fill_type="solid")
_BORDER = Border(*(Side(style='thin', color="D0D5DD"),) * 4)


def save_usg_rows(rows, filepath):
    """Сохраняет результат поиска USG (список строк-словарей) в .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Правила USG"

    headers = [title for _, title, _ in usg_parser.COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER

    for row in rows:
        ws.append(list(usg_parser.row_values(row)))
        fill = _DENY_FILL if row['action'] == 'deny' else _NAT_FILL if row['kind'] == 'nat-policy' else None
        for cell in ws[ws.max_row]:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill

    for idx, column in enumerate(ws.columns, 1):
        longest = max((len(str(c.value)) for c in column if c.value), default=0)
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, 50)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
